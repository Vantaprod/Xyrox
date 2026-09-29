"""
Xyrox — Excel Template Mapper (Phase 5.2)

Given a WorkbookInfo from Phase 5.1, build a TemplateMap that
identifies sections, field labels, and input cells — without any
hardcoded cell addresses or section names.

Master Command §15: build a generic engine, not a Lura-specific hack.
Master Command §16: actual workbook drives the map, not the image.

Design rules:
- No hardcoded cell addresses.
- No hardcoded section names.
- Deterministic: same workbook → same map.
- Report uncertainty honestly via confidence levels.
- Never invent data — if a label has no input, flag it.
"""

from dataclasses import dataclass, field, asdict
from typing import Optional

from openpyxl.utils import (
    coordinate_to_tuple,
    get_column_letter,
)

from app.backend.excel.inspector import WorkbookInfo, SheetInfo, CellInfo


class TemplateMappingError(Exception):
    """Raised when a workbook cannot be mapped."""
    pass


# ============================================================
# DATA TYPES
# ============================================================
@dataclass
class FieldMap:
    label: str
    label_cell: str
    input_cell: Optional[str]
    input_type: str          # text | integer | number | currency | date | percent | boolean
    number_format: Optional[str]
    label_is_bold: bool
    input_is_merged: bool
    confidence: str          # high | medium | low


@dataclass
class SectionMap:
    name: str
    header_cell: str
    start_row: int
    end_row: int
    fields: list[FieldMap] = field(default_factory=list)


@dataclass
class TemplateMap:
    file_path: str
    sheet_name: str
    sections: list[SectionMap] = field(default_factory=list)
    orphan_labels: list[str] = field(default_factory=list)
    blank_input_areas: list[str] = field(default_factory=list)


# ============================================================
# KEYWORD TABLES
# ============================================================
MIN_SECTION_TEXT_LEN = 3
WIDE_MERGE_COLS = 3

_INT_KEYWORDS = {
    "count", "number", "videos", "qty", "quantity", "units",
    "participants", "leads", "views", "reach", "followers",
    "customers", "cases", "sessions", "calls",
}
_DATE_KEYWORDS = {"date", "when", "day of", "month", "week", "year"}
_CURRENCY_KEYWORDS = {
    "amount", "revenue", "cost", "spend", "spent", "price",
    "value", "fund", "budget", "naira", "dollar", "income",
}
_PERCENT_KEYWORDS = {"percent", "percentage", "rate", "ratio", "%"}
_BOOL_KEYWORDS = {
    "completed?", "done?", "approved?", "yes/no", "flag",
}


# ============================================================
# HEURISTICS
# ============================================================
def _parse_merge_cols(merge_range: str) -> int:
    """
    Given a merge range like 'A2:C2', return the number of columns spanned.
    Returns 0 if it cannot be parsed.
    """
    try:
        left, right = merge_range.split(":")
        _, lcol = coordinate_to_tuple(left)
        _, rcol = coordinate_to_tuple(right)
        return rcol - lcol + 1
    except Exception:
        return 0


def _is_wide_merge(cell: CellInfo) -> bool:
    if not cell.is_merged or not cell.merge_range:
        return False
    return _parse_merge_cols(cell.merge_range) >= WIDE_MERGE_COLS


def _looks_like_section_header(cell: CellInfo) -> bool:
    """Heuristic: is this cell a section header?"""
    if cell.is_formula:
        return False
    v = cell.value
    if not isinstance(v, str):
        return False
    text = v.strip()
    if len(text) < MIN_SECTION_TEXT_LEN:
        return False

    wide = _is_wide_merge(cell)
    big_font = cell.font_size is not None and cell.font_size >= 12
    bold = cell.font_bold
    all_caps = text == text.upper() and any(ch.isalpha() for ch in text)

    if bold and (big_font or wide):
        return True
    if wide and all_caps:
        return True
    upper = text.upper()
    if upper.startswith("SECTION") or upper.endswith(" SECTION"):
        return True
    return False


def _infer_input_type(label_text: str, input_cell: Optional[CellInfo]) -> str:
    """
    Infer the expected type for a field.
    Priority: number_format on the input cell, then label keywords.
    """
    if input_cell is not None and input_cell.number_format:
        fmt = input_cell.number_format.lower()
        if any(t in fmt for t in ("yy", "dd", "mmm", "yyyy")):
            return "date"
        if any(t in fmt for t in ("$", "₦", "£", "€")):
            return "currency"
        if "%" in fmt:
            return "percent"
        if "." in fmt and "0" in fmt:
            return "number"
        if fmt in ("0", "#", "0.0", "0.00", "#,##0"):
            return "integer"

    text = label_text.lower()
    for k in _CURRENCY_KEYWORDS:
        if k in text:
            return "currency"
    for k in _PERCENT_KEYWORDS:
        if k in text:
            return "percent"
    for k in _DATE_KEYWORDS:
        if k in text:
            return "date"
    for k in _BOOL_KEYWORDS:
        if k in text:
            return "boolean"
    for k in _INT_KEYWORDS:
        if k in text:
            return "integer"
    return "text"


def _try_map_field(
    cell: CellInfo,
    cells_by_coord: dict[str, CellInfo],
) -> Optional[FieldMap]:
    """
    If this cell looks like a field label with an adjacent input cell,
    return a FieldMap. Otherwise return None.

    Tries the right neighbor first, then the cell below.
    """
    if cell.is_formula:
        return None
    v = cell.value
    if not isinstance(v, str):
        return None
    text = v.strip()
    if len(text) < 2 or len(text) > 200:
        return None

    # Skip cells that look like section headers themselves
    if _looks_like_section_header(cell):
        return None

    right_coord = f"{get_column_letter(cell.column + 1)}{cell.row}"
    below_coord = f"{get_column_letter(cell.column)}{cell.row + 1}"

    for coord in (right_coord, below_coord):
        neighbor = cells_by_coord.get(coord)

        if neighbor is not None:
            # Occupied cell → not an input slot
            if neighbor.is_formula or neighbor.value is not None:
                continue
            # Empty cell with border/fill/number_format → strong input
            if (
                neighbor.border
                or neighbor.fill_color
                or (neighbor.number_format and neighbor.number_format != "General")
            ):
                return FieldMap(
                    label=text,
                    label_cell=cell.coordinate,
                    input_cell=coord,
                    input_type=_infer_input_type(text, neighbor),
                    number_format=neighbor.number_format,
                    label_is_bold=bool(cell.font_bold),
                    input_is_merged=bool(neighbor.is_merged),
                    confidence="high",
                )
            # Empty cell with no styling → weaker input
            return FieldMap(
                label=text,
                label_cell=cell.coordinate,
                input_cell=coord,
                input_type=_infer_input_type(text, neighbor),
                number_format=neighbor.number_format,
                label_is_bold=bool(cell.font_bold),
                input_is_merged=bool(neighbor.is_merged),
                confidence="medium",
            )

        # Neighbor does not exist in the map (truly empty).
        # Only accept as input if the label looks intentional.
        if cell.font_bold or text.endswith(("?", ":")):
            return FieldMap(
                label=text,
                label_cell=cell.coordinate,
                input_cell=coord,
                input_type=_infer_input_type(text, None),
                number_format=None,
                label_is_bold=bool(cell.font_bold),
                input_is_merged=False,
                confidence="low",
            )
        return None

    return None


# ============================================================
# PUBLIC API
# ============================================================
def map_template(
    info: WorkbookInfo,
    sheet_name: Optional[str] = None,
) -> TemplateMap:
    """
    Build a TemplateMap from a WorkbookInfo.

    If sheet_name is None, uses the active sheet.
    Raises TemplateMappingError if the sheet cannot be found.
    """
    target = sheet_name or info.active_sheet
    if not target:
        raise TemplateMappingError("No sheet to map — workbook is empty.")

    sheet: Optional[SheetInfo] = next(
        (s for s in info.sheets if s.name == target), None
    )
    if sheet is None:
        raise TemplateMappingError(f"Sheet '{target}' not found.")

    cells = info.cells_by_sheet.get(target, [])
    if not cells:
        return TemplateMap(file_path=info.file_path, sheet_name=target)

    cells_by_coord: dict[str, CellInfo] = {c.coordinate: c for c in cells}

    # ---- Section header detection ----
    header_cells = [c for c in cells if _looks_like_section_header(c)]
    header_cells.sort(key=lambda c: (c.row, c.column))
    header_coords = {c.coordinate for c in header_cells}
    header_merge_ranges = {
        c.merge_range for c in header_cells if c.merge_range
    }

    def _is_inside_header(c: CellInfo) -> bool:
        if c.coordinate in header_coords:
            return True
        if c.merge_range and c.merge_range in header_merge_ranges:
            return True
        return False

    # ---- Build section boundaries ----
    sections: list[SectionMap] = []
    if header_cells:
        for i, h in enumerate(header_cells):
            start_row = h.row
            end_row = (
                header_cells[i + 1].row - 1
                if i + 1 < len(header_cells)
                else sheet.max_row
            )
            sections.append(SectionMap(
                name=str(h.value).strip(),
                header_cell=h.coordinate,
                start_row=start_row,
                end_row=end_row,
                fields=[],
            ))
    else:
        # No section headers found — one implicit section
        sections.append(SectionMap(
            name="(no section)",
            header_cell="A1",
            start_row=sheet.min_row,
            end_row=sheet.max_row,
            fields=[],
        ))

    # ---- Field detection ----
    orphan_labels: list[str] = []
    for c in cells:
        if _is_inside_header(c):
            continue
        if not isinstance(c.value, str):
            continue

        fm = _try_map_field(c, cells_by_coord)
        if fm is None:
            # If the cell is bold and has text, it may be an orphan label
            if c.font_bold and c.value and len(c.value.strip()) >= 2:
                orphan_labels.append(c.coordinate)
            continue

        # Assign to the section containing this label's row
        assigned = False
        for sec in sections:
            if sec.start_row <= c.row <= sec.end_row:
                sec.fields.append(fm)
                assigned = True
                break
        if not assigned and sections:
            sections[0].fields.append(fm)

    # ---- Sort sections and fields for deterministic output ----
    sections.sort(key=lambda s: s.start_row)
    for sec in sections:
        sec.fields.sort(
            key=lambda f: cells_by_coord[f.label_cell].row
            if f.label_cell in cells_by_coord
            else 0
        )

    # ---- Find unused blank input areas (styled empty cells) ----
    consumed_inputs = {
        f.input_cell
        for sec in sections
        for f in sec.fields
        if f.input_cell
    }
    blank_input_areas: list[str] = []
    for c in cells:
        if c.value is not None or c.is_formula:
            continue
        if c.coordinate in consumed_inputs:
            continue
        if (
            c.border
            or c.fill_color
            or (c.number_format and c.number_format != "General")
        ):
            blank_input_areas.append(c.coordinate)

    return TemplateMap(
        file_path=info.file_path,
        sheet_name=target,
        sections=sections,
        orphan_labels=orphan_labels,
        blank_input_areas=blank_input_areas,
    )


def template_map_to_dict(tm: TemplateMap) -> dict:
    """Convert TemplateMap into a JSON-serializable dict."""
    return {
        "file_path": tm.file_path,
        "sheet_name": tm.sheet_name,
        "sections": [
            {
                "name": s.name,
                "header_cell": s.header_cell,
                "start_row": s.start_row,
                "end_row": s.end_row,
                "fields": [asdict(f) for f in s.fields],
            }
            for s in tm.sections
        ],
        "orphan_labels": tm.orphan_labels,
        "blank_input_areas": tm.blank_input_areas,
    }