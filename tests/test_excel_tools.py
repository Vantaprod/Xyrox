"""
Xyrox — Excel Tool Tests (Phase 5.4)

Verifies the three AI-facing Excel tools are registered, callable,
and return JSON-serializable results. Uses real xlsx files.
"""

import json
import pytest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Border, Side

from app.backend.tools import build_default_registry
from app.backend.tools.base import ToolValidationError


def _box() -> Border:
    s = Side()
    return Border(left=s, right=s, top=s, bottom=s)


@pytest.fixture
def template_xlsx(tmp_path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "BRANDING & CONTENT"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Installation Videos"
    ws["B3"].border = _box()
    ws["B3"].number_format = "0"
    ws["A4"] = "Total Hours"
    ws["B4"].border = _box()
    ws["B4"].number_format = "0.00"
    ws["B6"] = "=B3*2"  # formula
    p = tmp_path / "template.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def output_dir(tmp_path) -> Path:
    d = tmp_path / "exports"
    d.mkdir()
    return d


# ============================================================
# REGISTRATION
# ============================================================
def test_excel_tools_registered():
    r = build_default_registry()
    assert "inspect_excel" in r
    assert "map_excel_template" in r
    assert "fill_excel_template" in r


def test_excel_tools_in_schema():
    r = build_default_registry()
    names = {t["function"]["name"] for t in r.to_ai_schema()}
    assert {"inspect_excel", "map_excel_template", "fill_excel_template"} <= names


# ============================================================
# inspect_excel
# ============================================================
def test_inspect_excel_tool(template_xlsx):
    r = build_default_registry()
    result = r.execute("inspect_excel", {"path": str(template_xlsx)})
    assert "sheet_names" in result
    assert "Report" in result["sheet_names"]
    assert "cells_by_sheet" in result
    assert len(result["cells_by_sheet"]["Report"]) > 0
    # JSON serializable
    json.dumps(result)


def test_inspect_excel_tool_missing_file(tmp_path):
    r = build_default_registry()
    with pytest.raises(ToolValidationError):
        r.execute("inspect_excel", {"path": str(tmp_path / "nope.xlsx")})


def test_inspect_excel_tool_empty_path():
    r = build_default_registry()
    with pytest.raises(ToolValidationError):
        r.execute("inspect_excel", {"path": ""})


def test_inspect_excel_tool_without_cells(template_xlsx):
    r = build_default_registry()
    result = r.execute("inspect_excel", {
        "path": str(template_xlsx),
        "include_cells": False,
    })
    assert result["cells_by_sheet"]["Report"] == []


# ============================================================
# map_excel_template
# ============================================================
def test_map_excel_template_tool(template_xlsx):
    r = build_default_registry()
    result = r.execute("map_excel_template", {"path": str(template_xlsx)})
    assert "sections" in result
    assert len(result["sections"]) == 1
    sec = result["sections"][0]
    assert sec["name"] == "BRANDING & CONTENT"
    labels = {f["label"] for f in sec["fields"]}
    assert "Installation Videos" in labels
    assert "Total Hours" in labels
    json.dumps(result)


def test_map_excel_template_tool_missing_file(tmp_path):
    r = build_default_registry()
    with pytest.raises(ToolValidationError):
        r.execute("map_excel_template", {"path": str(tmp_path / "nope.xlsx")})


def test_map_excel_template_tool_bad_sheet(template_xlsx):
    r = build_default_registry()
    with pytest.raises(ToolValidationError):
        r.execute("map_excel_template", {
            "path": str(template_xlsx),
            "sheet_name": "NopeSheet",
        })


# ============================================================
# fill_excel_template
# ============================================================
def test_fill_excel_template_tool(template_xlsx, output_dir):
    r = build_default_registry()
    result = r.execute("fill_excel_template", {
        "source_path": str(template_xlsx),
        "cell_values": {"B3": 5, "B4": 12.5},
        "output_dir": str(output_dir),
    })
    assert result["ok"] is True
    assert result["written"]["B3"] == 5
    assert Path(result["output_path"]).exists()
    json.dumps(result)


def test_fill_excel_template_preserves_original(template_xlsx, output_dir):
    before = template_xlsx.read_bytes()
    r = build_default_registry()
    r.execute("fill_excel_template", {
        "source_path": str(template_xlsx),
        "cell_values": {"B3": 42},
        "output_dir": str(output_dir),
    })
    after = template_xlsx.read_bytes()
    assert before == after


def test_fill_excel_template_refuses_formula(template_xlsx, output_dir):
    r = build_default_registry()
    result = r.execute("fill_excel_template", {
        "source_path": str(template_xlsx),
        "cell_values": {"B6": 999},
        "output_dir": str(output_dir),
    })
    assert result["ok"] is False
    assert "B6" in result["skipped"]
    # Formula preserved
    wb = load_workbook(result["output_path"])
    try:
        assert wb["Report"]["B6"].value == "=B3*2"
    finally:
        wb.close()


def test_fill_excel_template_missing_source(tmp_path, output_dir):
    r = build_default_registry()
    with pytest.raises(ToolValidationError):
        r.execute("fill_excel_template", {
            "source_path": str(tmp_path / "nope.xlsx"),
            "cell_values": {"A1": 1},
            "output_dir": str(output_dir),
        })


def test_fill_excel_template_empty_values(template_xlsx, output_dir):
    r = build_default_registry()
    with pytest.raises(ToolValidationError):
        r.execute("fill_excel_template", {
            "source_path": str(template_xlsx),
            "cell_values": {},
            "output_dir": str(output_dir),
        })


def test_fill_excel_template_invalid_coordinate(template_xlsx, output_dir):
    r = build_default_registry()
    with pytest.raises(ToolValidationError):
        r.execute("fill_excel_template", {
            "source_path": str(template_xlsx),
            "cell_values": {"NOT_A_CELL": 5},
            "output_dir": str(output_dir),
        })


# ============================================================
# MASTER COMMAND §15-17 — END-TO-END
# ============================================================
def test_acceptance_inspect_map_fill_verify(template_xlsx, output_dir):
    """
    The full Excel workflow a user can trigger via chat:
      1. inspect_excel  → understand structure
      2. map_excel_template → find cells
      3. fill_excel_template → write + verify
    """
    r = build_default_registry()

    # 1. Inspect
    info = r.execute("inspect_excel", {"path": str(template_xlsx)})
    assert "Report" in info["sheet_names"]

    # 2. Map
    tm = r.execute("map_excel_template", {"path": str(template_xlsx)})
    assert tm["sections"][0]["name"] == "BRANDING & CONTENT"

    # 3. Fill using the mapped cells
    fields = {f["label"]: f for f in tm["sections"][0]["fields"]}
    install_videos_cell = fields["Installation Videos"]["input_cell"]
    total_hours_cell = fields["Total Hours"]["input_cell"]

    fill = r.execute("fill_excel_template", {
        "source_path": str(template_xlsx),
        "cell_values": {
            install_videos_cell: 7,
            total_hours_cell: 3.5,
        },
        "output_dir": str(output_dir),
    })

    assert fill["ok"] is True
    assert fill["written"][install_videos_cell] == 7
    assert fill["written"][total_hours_cell] == 3.5

    # Verify by re-reading the output
    wb = load_workbook(fill["output_path"])
    try:
        assert wb["Report"][install_videos_cell].value == 7
        assert wb["Report"][total_hours_cell].value == 3.5
        # Original untouched
        assert wb["Report"]["B6"].value == "=B3*2"
    finally:
        wb.close()