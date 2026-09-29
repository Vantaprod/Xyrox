"""
Xyrox — Template Mapper Tests (Phase 5.2)

Builds real xlsx files and maps them. Proves section detection,
field detection, type inference, and honest reporting of ambiguity.
"""

import json

import pytest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side

from app.backend.excel.inspector import inspect_workbook
from app.backend.excel.template_mapper import (
    map_template,
    template_map_to_dict,
    TemplateMappingError,
)


# ============================================================
# HELPERS
# ============================================================
def _box() -> Border:
    s = Side()
    return Border(left=s, right=s, top=s, bottom=s)


# ============================================================
# FIXTURES
# ============================================================
@pytest.fixture
def simple_template(tmp_path) -> Path:
    """One section header + three fields with number formats."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    ws["A1"] = "BRANDING & CONTENT"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Installation Videos"
    ws["B3"].border = _box()
    ws["B3"].number_format = "0"

    ws["A4"] = "Product Demo Videos"
    ws["B4"].border = _box()
    ws["B4"].number_format = "0"

    ws["A5"] = "Follow-up Calls"
    ws["B5"].border = _box()
    ws["B5"].number_format = "0"

    path = tmp_path / "simple.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def multi_section_template(tmp_path) -> Path:
    """Two sections, each with two fields."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    ws["A1"] = "MARKET ENGAGEMENT"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Total Views"
    ws["B3"].border = _box()
    ws["B3"].number_format = "#,##0"
    ws["A4"] = "Total Reach"
    ws["B4"].border = _box()
    ws["B4"].number_format = "#,##0"

    ws["A6"] = "MARKETING FUND"
    ws["A6"].font = Font(bold=True, size=14)
    ws["A8"] = "Total Spent"
    ws["B8"].border = _box()
    ws["B8"].number_format = "₦#,##0.00"
    ws["A9"] = "Remaining Budget"
    ws["B9"].border = _box()
    ws["B9"].number_format = "₦#,##0.00"

    path = tmp_path / "multi.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def empty_template(tmp_path) -> Path:
    wb = Workbook()
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def date_template(tmp_path) -> Path:
    """Template with a date field."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "REPORTING PERIOD"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Start Date"
    ws["B3"].border = _box()
    ws["B3"].number_format = "yyyy-mm-dd"
    ws["A4"] = "End Date"
    ws["B4"].border = _box()
    ws["B4"].number_format = "yyyy-mm-dd"
    path = tmp_path / "dates.xlsx"
    wb.save(path)
    return path


# ============================================================
# BASIC MAPPING
# ============================================================
def test_map_simple_template(simple_template):
    info = inspect_workbook(str(simple_template))
    tm = map_template(info, "Report")

    assert tm.sheet_name == "Report"
    assert len(tm.sections) == 1
    sec = tm.sections[0]
    assert sec.name == "BRANDING & CONTENT"
    assert sec.header_cell == "A1"
    assert len(sec.fields) == 3

    labels = {f.label for f in sec.fields}
    assert labels == {
        "Installation Videos",
        "Product Demo Videos",
        "Follow-up Calls",
    }

    for f in sec.fields:
        assert f.input_cell is not None
        assert f.input_cell.startswith("B")
        assert f.input_type == "integer"
        assert f.confidence == "high"


def test_map_multi_section_template(multi_section_template):
    info = inspect_workbook(str(multi_section_template))
    tm = map_template(info, "Report")

    assert len(tm.sections) == 2
    names = [s.name for s in tm.sections]
    assert names == ["MARKET ENGAGEMENT", "MARKETING FUND"]

    s1 = tm.sections[0]
    assert len(s1.fields) == 2
    assert {f.label for f in s1.fields} == {"Total Views", "Total Reach"}
    for f in s1.fields:
        assert f.input_type == "integer"

    s2 = tm.sections[1]
    assert len(s2.fields) == 2
    for f in s2.fields:
        # currency inferred from ₦ symbol in number_format
        assert f.input_type == "currency"


def test_map_date_template(date_template):
    info = inspect_workbook(str(date_template))
    tm = map_template(info, "Report")
    assert len(tm.sections) == 1
    for f in tm.sections[0].fields:
        assert f.input_type == "date"


# ============================================================
# DEFAULTS / EMPTY
# ============================================================
def test_map_empty_workbook(empty_template):
    info = inspect_workbook(str(empty_template))
    tm = map_template(info, info.active_sheet)
    # No cells, empty sections
    assert tm.sections == []


def test_map_nonexistent_sheet(simple_template):
    info = inspect_workbook(str(simple_template))
    with pytest.raises(TemplateMappingError):
        map_template(info, "DoesNotExist")


def test_map_uses_active_sheet_by_default(simple_template):
    info = inspect_workbook(str(simple_template))
    tm = map_template(info)
    assert tm.sheet_name == info.active_sheet


# ============================================================
# SECTION DETECTION
# ============================================================
def test_section_detected_by_bold_large_font(simple_template):
    info = inspect_workbook(str(simple_template))
    tm = map_template(info)
    assert any(s.name == "BRANDING & CONTENT" for s in tm.sections)


def test_section_detected_by_wide_merge(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "TRAINING & ACTIVATION"
    ws.merge_cells("A1:F1")
    ws["A1"].font = Font(bold=True)

    ws["A3"] = "Sessions Held"
    ws["B3"].border = _box()
    ws["B3"].number_format = "0"

    path = tmp_path / "merged.xlsx"
    wb.save(path)

    info = inspect_workbook(str(path))
    tm = map_template(info)
    assert any(s.name == "TRAINING & ACTIVATION" for s in tm.sections)


# ============================================================
# TYPE INFERENCE
# ============================================================
def test_type_inference_from_number_format(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "FINANCE"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Amount"
    ws["B3"].border = _box()
    ws["B3"].number_format = "$#,##0.00"

    ws["A4"] = "Rate"
    ws["B4"].border = _box()
    ws["B4"].number_format = "0.00%"

    ws["A5"] = "Count"
    ws["B5"].border = _box()
    ws["B5"].number_format = "0"

    path = tmp_path / "types.xlsx"
    wb.save(path)

    info = inspect_workbook(str(path))
    tm = map_template(info)
    types = {f.label: f.input_type for f in tm.sections[0].fields}
    assert types["Amount"] == "currency"
    assert types["Rate"] == "percent"
    assert types["Count"] == "integer"


def test_type_inference_from_label_keywords(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "GENERAL"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Total Videos Produced"
    ws["B3"].border = _box()
    path = tmp_path / "labels.xlsx"
    wb.save(path)

    info = inspect_workbook(str(path))
    tm = map_template(info)
    assert tm.sections[0].fields[0].input_type == "integer"


# ============================================================
# HONESTY / UNCERTAINTY
# ============================================================
def test_orphan_label_reported(tmp_path):
    """A bold label with an occupied right cell → orphan."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "SECTION ONE"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Lonely Label"
    ws["A3"].font = Font(bold=True)
    ws["B3"] = "already here"   # occupied, so not an input

    path = tmp_path / "orphan.xlsx"
    wb.save(path)

    info = inspect_workbook(str(path))
    tm = map_template(info)
    # The label should NOT be a field, but should be reported as orphan
    # only if no input found. Below-neighbor is also empty though...
    # Our heuristic will use below if right is occupied and empty below.
    # To ensure the orphan case, put something below too.
    # Rewriting the test to force the orphan case:
    ws["A4"] = "already here too"
    wb.save(path)
    info = inspect_workbook(str(path))
    tm = map_template(info)
    assert "A3" in tm.orphan_labels


# ============================================================
# DETERMINISM
# ============================================================
def test_map_is_deterministic(simple_template):
    info1 = inspect_workbook(str(simple_template))
    info2 = inspect_workbook(str(simple_template))
    tm1 = map_template(info1)
    tm2 = map_template(info2)

    d1 = template_map_to_dict(tm1)
    d2 = template_map_to_dict(tm2)
    assert d1 == d2


# ============================================================
# SERIALIZATION
# ============================================================
def test_template_map_is_json_serializable(simple_template):
    info = inspect_workbook(str(simple_template))
    tm = map_template(info)
    d = template_map_to_dict(tm)
    js = json.dumps(d)
    assert "BRANDING & CONTENT" in js
    assert "Installation Videos" in js


# ============================================================
# MASTER COMMAND §15 — ADAPTIVE, NOT HARDCODED
# ============================================================
def test_mapper_is_adaptive_not_hardcoded(tmp_path):
    """
    Two different templates with different section names must both
    map correctly. Proves there is no hardcoded Lura logic.
    """
    # Template A
    wb = Workbook()
    ws = wb.active
    ws.title = "A"
    ws["A1"] = "SALES & GROWTH"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "New Customers"
    ws["B3"].border = _box()
    ws["B3"].number_format = "0"
    path_a = tmp_path / "a.xlsx"
    wb.save(path_a)

    # Template B
    wb = Workbook()
    ws = wb.active
    ws.title = "B"
    ws["A1"] = "OPERATIONS"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Tickets Resolved"
    ws["B3"].border = _box()
    ws["B3"].number_format = "0"
    path_b = tmp_path / "b.xlsx"
    wb.save(path_b)

    tm_a = map_template(inspect_workbook(str(path_a)))
    tm_b = map_template(inspect_workbook(str(path_b)))

    assert tm_a.sections[0].name == "SALES & GROWTH"
    assert tm_b.sections[0].name == "OPERATIONS"
    assert tm_a.sections[0].fields[0].label == "New Customers"
    assert tm_b.sections[0].fields[0].label == "Tickets Resolved"