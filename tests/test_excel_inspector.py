"""
Xyrox — Excel Inspector Tests (Phase 5.1)

Builds real xlsx files with openpyxl and inspects them.
Proves the inspector reads values, formulas, merges, and
formatting without modifying the file.
"""

import pytest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.backend.excel.inspector import (
    inspect_workbook,
    workbook_info_to_dict,
    ExcelInspectionError,
)


# ============================================================
# FIXTURES
# ============================================================
@pytest.fixture
def simple_xlsx(tmp_path) -> Path:
    """A workbook with 2 sheets, values, a formula, and formatting."""
    wb = Workbook()

    # Sheet 1: report
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Branding & Content"
    ws["A1"].font = Font(bold=True, size=14)
    ws["F8"] = "Installation Videos"
    ws["G8"] = 5
    ws["G8"].number_format = "0"
    ws["G9"] = "=G8*2"
    ws["A2"] = "Section Header"
    ws["A2"].fill = PatternFill(start_color="FFCC00", end_color="FFCC00", fill_type="solid")
    ws.merge_cells("A2:C2")
    ws["A2"].alignment = Alignment(horizontal="center")

    # Hidden row and column
    ws.row_dimensions[5].hidden = True
    ws.column_dimensions["D"].hidden = True

    # Sheet 2: notes
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Just a note"

    path = tmp_path / "test_report.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def empty_xlsx(tmp_path) -> Path:
    wb = Workbook()
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    return path


# ============================================================
# INSPECTION
# ============================================================
def test_inspect_simple_workbook(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    assert info.sheet_names == ["Report", "Notes"]
    assert info.active_sheet == "Report"
    assert info.has_macros is False
    assert info.file_size_bytes > 0


def test_inspect_reads_sheet_metadata(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    report = next(s for s in info.sheets if s.name == "Report")
    assert report.index == 0
    assert report.state == "visible"
    assert report.max_row >= 9
    assert report.max_column >= 7


def test_inspect_detects_merged_range(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    report = next(s for s in info.sheets if s.name == "Report")
    assert "A2:C2" in report.merged_ranges


def test_inspect_detects_hidden_row_and_column(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    report = next(s for s in info.sheets if s.name == "Report")
    assert 5 in report.hidden_rows
    assert "D" in report.hidden_columns


def test_inspect_reads_values(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    report_cells = info.cells_by_sheet["Report"]
    by_coord = {c.coordinate: c for c in report_cells}

    assert "A1" in by_coord
    assert by_coord["A1"].value == "Branding & Content"
    assert by_coord["F8"].value == "Installation Videos"
    assert by_coord["G8"].value == 5
    assert by_coord["G8"].number_format == "0"


def test_inspect_detects_formula(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    report_cells = info.cells_by_sheet["Report"]
    by_coord = {c.coordinate: c for c in report_cells}
    assert "G9" in by_coord
    assert by_coord["G9"].is_formula is True
    assert by_coord["G9"].formula == "=G8*2"
    # Value is None for formulas when data_only=False
    assert by_coord["G9"].value is None


def test_inspect_detects_bold_font(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    by_coord = {c.coordinate: c for c in info.cells_by_sheet["Report"]}
    assert by_coord["A1"].font_bold is True


def test_inspect_detects_fill_color(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    by_coord = {c.coordinate: c for c in info.cells_by_sheet["Report"]}
    fill = by_coord["A2"].fill_color
    assert fill is not None
    # openpyxl normalizes to 8-char hex (alpha+RGB)
    assert "FFCC00" in fill.upper()


def test_inspect_marks_merged_cells(simple_xlsx):
    info = inspect_workbook(str(simple_xlsx))
    by_coord = {c.coordinate: c for c in info.cells_by_sheet["Report"]}
    # Top-left of merge
    assert by_coord["A2"].is_merged is True
    assert by_coord["A2"].merge_range == "A2:C2"
    # Middle of merge
    assert by_coord["B2"].is_merged is True
    assert by_coord["B2"].merge_range == "A2:C2"


def test_inspect_without_cells(simple_xlsx):
    """include_cells=False skips cell reading but keeps sheet metadata."""
    info = inspect_workbook(str(simple_xlsx), include_cells=False)
    report = next(s for s in info.sheets if s.name == "Report")
    assert "A2:C2" in report.merged_ranges
    # No cell info
    assert info.cells_by_sheet["Report"] == []


def test_inspect_empty_workbook(empty_xlsx):
    info = inspect_workbook(str(empty_xlsx))
    assert len(info.sheet_names) >= 1
    # Default Sheet has minimal cells
    for sheet in info.sheets:
        assert sheet.name


# ============================================================
# ERROR HANDLING
# ============================================================
def test_inspect_missing_file(tmp_path):
    with pytest.raises(ExcelInspectionError):
        inspect_workbook(str(tmp_path / "nope.xlsx"))


def test_inspect_wrong_extension(tmp_path):
    p = tmp_path / "not_xlsx.txt"
    p.write_text("hello")
    with pytest.raises(ExcelInspectionError):
        inspect_workbook(str(p))


def test_inspect_corrupt_file(tmp_path):
    p = tmp_path / "corrupt.xlsx"
    p.write_bytes(b"not a real xlsx file")
    with pytest.raises(ExcelInspectionError):
        inspect_workbook(str(p))


# ============================================================
# SERIALIZATION
# ============================================================
def test_workbook_info_to_dict_is_json_serializable(simple_xlsx):
    import json
    info = inspect_workbook(str(simple_xlsx))
    d = workbook_info_to_dict(info)
    # Must not raise
    js = json.dumps(d)
    assert "Report" in js
    assert "A2:C2" in js


def test_inspection_does_not_modify_file(simple_xlsx):
    before = simple_xlsx.read_bytes()
    inspect_workbook(str(simple_xlsx))
    after = simple_xlsx.read_bytes()
    assert before == after, "Inspector modified the source file."